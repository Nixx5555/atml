"""
make_ppt.py - builds outputs/RikoNet_Replication_Presentation.pptx from outputs/results.json.
Charts are native PowerPoint charts (editable); speaker notes explain each slide.
Run AFTER main.py.  python make_ppt.py
"""
import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import config as C
from report_common import KG, OURS, PAPER, PAPER_TABLE4, facts, fmt, row, seeds_text, short

f = facts()
R, F = f["R"], f["F"]

INK, INK2, MUTED = "1F1B3A", "3D3A55", "6B6880"
BLUE, ORANGE, AQUA, YELLOW = "2A78D6", "EB6834", "1BAF7A", "EDA100"
GREY, TINT, LIGHT, WHITE = "A9A8A2", "EAF1FB", "F4F4F7", "FFFFFF"
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, s, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
        r.font.color.rgb = rgb(color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.28))))
        pPr.set("indent", str(-int(Inches(0.28))))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")
        p.space_after = Pt(gap)
        lead, rest = it if isinstance(it, tuple) else ("", it)
        if lead:
            r = p.add_run()
            r.text = lead
            r.font.bold, r.font.size, r.font.name = True, Pt(size), BODY
            r.font.color.rgb = rgb(color)
        r = p.add_run()
        r.text = rest
        r.font.size, r.font.name = Pt(size), BODY
        r.font.color.rgb = rgb(color)
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def label_box(slide, x, y, w, h, s, fill, color=INK, size=13, bold=False):
    b = box(slide, x, y, w, h, fill)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size, r.font.name, r.font.bold = Pt(size if i == 0 else size - 2), BODY, bold and i == 0
        r.font.color.rgb = rgb(color)
    return b


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def badge(slide, x, y, n, fill=BLUE, d=0.5):
    b = box(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size, r.font.bold, r.font.name = Pt(16), True, BODY
    r.font.color.rgb = rgb(WHITE)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def title(slide, s, sub=None):
    text(slide, 0.6, 0.38, W - 1.2, 0.8, s, size=32, bold=True, font=HEAD, color=INK)
    if sub:
        text(slide, 0.6, 1.12, W - 1.2, 0.45, sub, size=15, color=MUTED, italic=True)


def footer(slide, n):
    text(slide, 0.6, 7.02, 8, 0.3, "RikoNet replication · ATML", size=10, color=MUTED)
    text(slide, W - 1.6, 7.02, 1.0, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def table(slide, x, y, w, df, col_w=None, size=12, highlight=(), row_h=0.36):
    rows, cols = len(df) + 1, len(df.columns)
    t = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows)).table
    for j, c in enumerate(df.columns):
        if col_w:
            t.columns[j].width = Inches(col_w[j])
    for i in range(rows):
        t.rows[i].height = Inches(row_h)
        for j in range(cols):
            cell = t.cell(i, j)
            v = df.columns[j] if i == 0 else df.iloc[i - 1, j]
            cell.text = fmt(v) if isinstance(v, float) else str(v)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            for r in p.runs:
                r.font.size, r.font.name = Pt(size), BODY
                r.font.bold = i == 0 or (i - 1) in highlight
                r.font.color.rgb = rgb(WHITE if i == 0 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(INK if i == 0 else (TINT if (i - 1) in highlight else
                                                               (WHITE if i % 2 else LIGHT)))
    return t


def bar_chart(slide, x, y, w, h, labels, values, colors, number_format="0.000", title_=None,
              axis_max=None):
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("value", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size, ch.font.name = Pt(12), BODY
    ch.font.color.rgb = rgb(INK2)
    if title_:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_
        tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
        tr.font.size, tr.font.bold, tr.font.name = Pt(14), True, BODY
        tr.font.color.rgb = rgb(INK)
    else:
        ch.has_title = False
    pl = ch.plots[0]
    pl.gap_width = 45
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = number_format, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(11)
    dl.font.color.rgb = rgb(INK)
    for i, c in enumerate(colors):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E6E5E0")
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(10)
    va.minimum_scale = 0
    if axis_max:
        va.maximum_scale = axis_max
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb("D0CFCA")
    ca.tick_labels.font.size = Pt(11)
    ca.reverse_order = True
    return ch


def stat(slide, x, y, w, value, label, color=BLUE, size=48):
    text(slide, x, y, w, 0.95, value, size=size, bold=True, color=color, font=HEAD)
    text(slide, x, y + 0.95, w, 0.7, label, size=13, color=MUTED)


n = [0]


def new(dark=False):
    s = prs.slides.add_slide(BLANK)
    n[0] += 1
    if dark:
        bg(s, INK)
    else:
        bg(s, WHITE)
        footer(s, n[0])
    return s


e1o, e1b = f["e1_ours"], f["e1_best_baseline"]
cfg = R["config"]

# =================================================================== 1 title
s = new(dark=True)
text(s, 0.8, 1.0, 11, 0.5, "ADVANCED TOPICS IN MACHINE LEARNING · PAPER REPLICATION", size=14,
     color="9FB8E8", bold=True)
text(s, 0.8, 1.7, 11.5, 1.4, "Replicating RikoNet", size=54, bold=True, color=WHITE, font=HEAD)
text(s, 0.8, 2.85, 11.5, 1.0, "A deep-autoencoder + spectral-clustering hybrid recommender for anime",
     size=24, color="CADCFC")
text(s, 0.8, 4.4, 11.5, 0.9, [f"Soni, Thakuria, Nath, Das & Boro · {PAPER['venue']}",
                              f"DOI {PAPER['doi']}"], size=14, color="CADCFC")
text(s, 0.8, 5.9, 11.5, 0.5, "Presented by: ______________________     Supervisor: ______________________",
     size=14, color=WHITE)
notes(s, "Introduce yourself and the paper. RikoNet was published in 2023 in Springer's Multimedia Tools and "
         "Applications, a subscription journal. The goal of this project: re-implement the method from the "
         "paper alone and check whether its reported results hold.")

# =================================================================== 2 at a glance
s = new()
title(s, "The answer up front", "What we rebuilt, and whether the paper's claim held up")
stat(s, 0.7, 1.9, 3.6, "0.591", "RMSE on MovieLens-100K reported in the paper", ORANGE)
stat(s, 4.7, 1.9, 3.6, fmt(e1o.RMSE), f"our RikoNet RMSE, same data and protocol ({seeds_text(e1o.runs)})", BLUE)
stat(s, 8.7, 1.9, 3.9, fmt(e1b.RMSE), f"best baseline we trained: {short(e1b.model)}", GREY)
bullets(s, 0.7, 4.1, 12, 2.6, [
    ("Rebuilt end to end: ", "masked deep autoencoder, 3-stage embeddings, spectral clustering, "
                             "opposite clusters, hybrid filter."),
    ("Not reproduced: ", f"the headline RMSE. Ours is {fmt(f['e1_gap_to_paper'])} higher, and "
                         f"{f['e1_vs_baseline']} a tuned classical baseline."),
    ("Type: ", "simplified replication. The authors' anime data and code are not public, so we used the "
               "paper's public benchmark plus an alternative public anime dataset."),
], size=17)
notes(s, "Give the conclusion first so the audience knows where the talk is going. The three numbers: what the "
         "paper reports, what we got with the same dataset and protocol, and what a simple baseline gets.")

# =================================================================== 3 verification
s = new()
title(s, "Paper selection was verified before starting")
cards = [
    ("Subscription only", "Springer page metadata: Open Access = N, isAccessibleForFree = false. "
                          "(A preprint exists on arXiv 2106.12970.)"),
    ("Authors' code: gone", "github.com/NilutpalNath/RikoNet returns 404 and is not in the author's "
                            "public repository list."),
    ("Anime data: private", "Crawled from MyAnimeList via the Jikan API, never released, no "
                            "reconstruction instructions."),
    ("Benchmark data: public", "Table 4 uses MovieLens-100K, which we downloaded from GroupLens, so "
                               "that experiment can be reproduced exactly."),
]
for i, (h_, b_) in enumerate(cards):
    x = 0.6 + (i % 2) * 6.15
    y = 1.65 + (i // 2) * 2.55
    box(s, x, y, 5.95, 2.3, LIGHT)
    badge(s, x + 0.3, y + 0.3, i + 1, BLUE if i != 2 else ORANGE)
    text(s, x + 1.0, y + 0.28, 4.7, 0.5, h_, size=20, bold=True)
    text(s, x + 1.0, y + 0.85, 4.7, 1.4, b_, size=15, color=INK2)
notes(s, "Faculty asked for a closed-access Springer paper with public data. The journal version is "
         "subscription-only; mention honestly that an arXiv preprint exists. The anime dataset is private, which "
         "is why this is a simplified replication.")

# =================================================================== 4 architecture
s = new()
title(s, "How RikoNet works", "Collaborative signal (top) + content signal (bottom) meet in the hybrid filter")
y1, y2, bw, bh = 1.95, 4.35, 2.05, 1.05
top = [("User ratings", ["User ratings", "+ age, gender"]), ("Primary AE", ["Primary", "autoencoder"]),
       ("Predicted", ["Predicted ratings", "for unrated titles (S)"])]
xs = [0.6, 3.15, 5.7]
for x, (_, lines) in zip(xs, top):
    label_box(s, x, y1, bw, bh, lines, TINT, size=14, bold=True)
bot = [["Top raters'", "ratings per anime"], ["AE-1  Rᵐ→Rᵈ", "+ 43 genres"],
       ["AE-3  → hybrid", "embedding (d)"], ["Spectral", "clustering (222)"], ["Opposite", "clusters (Alg. 1)"]]
xs2 = [0.6, 3.15, 5.7, 8.25, 10.8]
for x, lines in zip(xs2, bot):
    label_box(s, x, y2, bw, bh, lines, LIGHT, size=14, bold=True)
label_box(s, 8.25, y1, 2.05, bh, ["Hybrid filter", "S ∩ C  /  S − C"], BLUE, WHITE, 14, True)
label_box(s, 10.8, y1 - 0.35, 2.05, 0.8, ["Similar Anime", "top 10"], INK, WHITE, 13, True)
label_box(s, 10.8, y1 + 0.6, 2.05, 0.8, ["You May Like", "top 10"], INK, WHITE, 13, True)
for a, b_ in zip(xs[:-1], xs[1:]):
    arrow(s, a + bw, y1 + bh / 2, b_, y1 + bh / 2)
arrow(s, 5.7 + bw, y1 + bh / 2, 8.25, y1 + bh / 2)
for a, b_ in zip(xs2[:-1], xs2[1:]):
    arrow(s, a + bw, y2 + bh / 2, b_, y2 + bh / 2)
arrow(s, 9.27, y2, 9.27, y1 + bh)
arrow(s, 10.3, y1 + 0.3, 10.8, y1 + 0.05)
arrow(s, 10.3, y1 + 0.75, 10.8, y1 + 1.0)
text(s, 0.6, 5.85, 12, 0.9, "C = union of the clusters of the user's last 3 rated titles: own cluster if liked, "
     "logically opposite cluster if disliked.", size=15, color=INK2)
notes(s, "Top row: the collaborative part. A user's ratings go through the autoencoder, which fills in ratings "
         "for everything unrated. Bottom row: the content part. Anime are embedded using ratings from heavy "
         "users plus genres, then clustered. The hybrid filter splits the predictions into titles inside the "
         "relevant clusters (Similar) and outside them (You May Like).")

# =================================================================== 5 autoencoder
s = new()
title(s, "The primary autoencoder", "Kuchaiev & Ginsburg (2017) architecture, as RikoNet specifies")
layers = [("n", 2.6), ("512", 2.0), ("512", 1.6), ("1024", 1.2), ("512", 1.6), ("512", 2.0), ("n", 2.6)]
for i, (lab, hgt) in enumerate(layers):
    x = 0.8 + i * 0.85
    col = BLUE if lab == "1024" else (INK if lab == "n" else "9FB8E8")
    box(s, x, 4.1 - hgt / 2, 0.55, hgt, col, radius=0.2)
    text(s, x - 0.2, 4.1 + hgt / 2 + 0.08, 0.95, 0.35, lab, size=12, color=INK2, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.85, 6.2, 0.9, ["K&G default sizes shown; code layer has dropout 0.8.",
                              f"Tuned on validation: {cfg['movielens_tuned']['hidden']} (MovieLens)"],
     size=12, color=MUTED)
bullets(s, 7.0, 1.75, 5.8, 5.0, [
    ("Input: ", "one user = ratings of all titles, 0 = unknown (+ gender and 5 age groups)."),
    ("Activations: ", "SELU in hidden layers, ReLU on the output (ratings are never negative)."),
    ("Masked MSE: ", "error is counted only where a real rating exists."),
    ("Dense re-feeding: ", "the dense output f(x) is fed back as a new input with target f(x)."),
    ("Prediction: ", "the reconstructed value at an unrated position is the predicted rating."),
], size=16)
notes(s, "The key idea of masked MSE: zeros mean 'unknown', so they must not be treated as real ratings of zero. "
         "Dense re-feeding is a data-augmentation trick from the NVIDIA paper.")

# =================================================================== 6 replication scope
s = new()
title(s, "Exact reproduction vs simplified replication")
scope = pd.DataFrame([
    ("Autoencoder: SELU/ReLU, masked MSE, re-feeding", "Faithful"),
    ("5 % per-user hold-out, MSE/RMSE", "Faithful"),
    ("MovieLens-100K benchmark (Table 4)", "Faithful, same public data"),
    ("Spectral clustering, 222 clusters; Algorithm 1; hybrid filter", "Faithful"),
    ("Layer sizes, dropout, optimiser, d, t, like threshold", "Not in paper: chosen on validation"),
    ("Self-tuning cluster count (Table 1)", "Approximated (eigengap)"),
    ("MyAnimeList anime data + demographics", "Substituted: Kaggle anime data"),
    ("'Last 3 rated' (needs timestamps)", "Approximated: 3 random known titles"),
    ("Baselines", "Re-trained on our split (paper copied them)"),
    ("20-person user study", "Not replicated"),
], columns=["Component", "Status"])
table(s, 0.6, 1.5, 12.1, scope, col_w=[7.2, 4.9], size=14, row_h=0.47, highlight=[4, 5, 6, 7])
notes(s, "Be explicit about this to the faculty: highlighted rows are where we had to deviate or guess, because "
         "the paper does not report the detail or the original resource is unavailable.")

# =================================================================== 7 data & protocol
s = new()
title(s, "Data and evaluation protocol")
ms = R["movielens_stats"]
box(s, 0.6, 1.6, 3.9, 2.55, LIGHT)
text(s, 0.85, 1.75, 3.5, 0.5, "MovieLens-100K", size=20, bold=True)
text(s, 0.85, 2.3, 3.5, 1.8, [f"{ms['users']} users · {ms['items']:,} movies", f"{ms['ratings']:,} ratings, 1-5",
                              "same data as paper Table 4", "has age + gender"], size=15, color=INK2)
if "anime_stats" in R:
    a = R["anime_stats"]
    box(s, 4.75, 1.6, 3.9, 2.55, LIGHT)
    text(s, 5.0, 1.75, 3.5, 0.5, "Anime (Kaggle)", size=20, bold=True)
    text(s, 5.0, 2.3, 3.5, 1.8, [f"{a['users']:,} users · {a['items']:,} titles", f"{a['ratings']:,} ratings, 1-10",
                                 f"{a['genres']} genres (same as paper)", "ALTERNATIVE dataset"], size=15, color=INK2)
segs = [("train 90 %", 0.9 * 3.9, BLUE), ("val 5 %", 0.05 * 3.9, AQUA), ("test 5 %", 0.05 * 3.9, ORANGE)]
x = 8.95
text(s, 8.9, 1.75, 4, 0.5, "Per-user split", size=20, bold=True)
for lab, wdt, col in segs:
    box(s, x, 2.45, wdt, 0.55, col, shape=MSO_SHAPE.RECTANGLE)
    x += wdt
text(s, 8.9, 3.1, 4, 1.0, ["train 90 %  ·  validation 5 %  ·  test 5 %",
                           "(same idea as the paper's 5 % hold-out)"], size=13, color=MUTED)
bullets(s, 0.6, 4.5, 12.1, 2.3, [
    ("No leakage: ", "test ratings are removed from every model's input; tuning uses validation only; "
                     "embeddings and clusters use training ratings only."),
    ("Fair baselines: ", "averages, bias model, user/item KNN, biased MF, popularity - all on the same split."),
    ("Metrics: ", "MSE, RMSE, MAE (paper) + Precision, Recall, NDCG, HitRate @10 (added)."),
], size=16)
notes(s, "Explain why splitting per user matters: every user keeps most of their history as input, and the hidden "
         "5 percent is what we try to predict. The validation split lets us tune without ever touching test.")

# =================================================================== 8 E1 results
s = new()
title(s, "Result 1: MovieLens-100K, paper protocol", "Test RMSE, lower is better")
e1 = F["E1"].sort_values("RMSE")
labels = [short(m) for m in e1.model] + ["RikoNet (paper)"]
vals = list(e1.RMSE) + [0.591]
cols = [BLUE if "RikoNet" in m else GREY for m in e1.model] + [ORANGE]
bar_chart(s, 0.5, 1.55, 7.9, 5.35, labels, vals, cols, axis_max=round(max(vals) * 1.15, 1))
bullets(s, 8.7, 1.8, 4.1, 5.0, [
    ("Ours: ", f"{fmt(e1o.RMSE)} ± {fmt(e1o.RMSE_std)} ({seeds_text(e1o.runs)})."),
    ("Best baseline: ", f"{short(e1b.model)}, {fmt(e1b.RMSE)}."),
    ("K&G default config: ", f"{fmt(row(F['E1'], KG).RMSE)}."),
    ("Paper: ", "0.591, which we could not approach."),
], size=16)
notes(s, "Blue bars are our RikoNet runs, grey bars are baselines trained by us on exactly the same split, and "
         "the orange bar is the number printed in the paper. All methods we trained sit in roughly the same band.")

# =================================================================== 9 Table 4 check
s = new()
title(s, "Why the gap? Looking closely at paper Table 4")
t4 = pd.DataFrame(PAPER_TABLE4, columns=["Method", "MSE", "RMSE"])
t4["√MSE"] = (t4.MSE ** 0.5).round(3)
t4["OK?"] = ["yes" if abs(a - b) < 0.002 else "NO" for a, b in zip(t4["√MSE"], t4.RMSE)]
table(s, 0.6, 1.45, 6.6, t4, col_w=[2.9, 0.9, 0.9, 0.9, 1.0], size=12, row_h=0.37, highlight=[2, 3, 12])
bullets(s, 7.6, 1.6, 5.2, 5.2, [
    ("Inconsistent rows: ", "for User/Item KNN-Cosine, RMSE is not √MSE."),
    ("Borrowed numbers: ", "competitors were copied from another paper, so they were not run on the 5 % protocol."),
    ("Missing details: ", "layer sizes, epochs, optimiser, d and t are not reported."),
    ("Our check: ", "re-training every baseline on one split gives a like-for-like comparison."),
], size=16)
notes(s, "This is an important critical-thinking slide for ATML: the published comparison mixes numbers from "
         "different sources, and two rows are arithmetically inconsistent.")

# =================================================================== 10 activation
if "E3_anime" in R or "E3_movielens" in R:
    s = new()
    key = "E3_anime" if "E3_anime" in R else "E3_movielens"
    ds = "anime (Kaggle)" if key == "E3_anime" else "MovieLens-100K"
    title(s, "Result 2: which activation works best?", f"Paper Fig. 3 replicated on {ds}")
    img = C.FIG_DIR / ("activation_study_anime.png" if key == "E3_anime" else "activation_study_movielens.png")
    s.shapes.add_picture(str(img), Inches(0.5), Inches(1.7), width=Inches(8.3))
    rows_ = sorted(R[key].items(), key=lambda kv: kv[1]["best_val_mse"])
    table(s, 9.1, 1.8, 3.7, pd.DataFrame([(a.upper(), round(v["best_val_mse"], 3)) for a, v in rows_],
                                         columns=["Activation", "Best val MSE"]),
          col_w=[1.7, 2.0], size=14, row_h=0.45, highlight=[0])
    verdict = "SELU is best: the paper's claim is confirmed." if f[key + "_best"] == "selu" else \
        f"{f[key + '_best'].upper()} is best here; SELU ranks {f[key + '_selu_rank']} of 4, so the paper's claim is not confirmed."
    text(s, 9.1, 4.3, 3.8, 2.0, verdict, size=16, color=INK, bold=True)
    notes(s, "Same network, only the activation function changes. The paper reports SELU as best; the table shows "
             "the lowest validation MSE each activation reached.")

# =================================================================== 11 anime rating + topk
if "E4" in F:
    s = new()
    title(s, "Result 3: anime data, rating prediction and top-10 lists")
    e4 = F["E4"].sort_values("RMSE")
    bar_chart(s, 0.5, 1.4, 6.3, 5.4, [short(m) for m in e4.model], list(e4.RMSE),
              [BLUE if "RikoNet" in m else GREY for m in e4.model], title_="Test RMSE (1-10 scale)",
              axis_max=round(max(e4.RMSE) * 1.18, 1))
    rk = F["E4_rank"].sort_values("NDCG@10", ascending=False)
    bar_chart(s, 6.9, 1.4, 6.0, 5.4, [short(m) for m in rk.model], list(rk["NDCG@10"]),
              [BLUE if "RikoNet" in m else GREY for m in rk.model], title_="NDCG@10 (higher is better)",
              axis_max=round(max(rk["NDCG@10"]) * 1.3, 2))
    notes(s, f"Left: rating error on held-out anime ratings. Right: quality of the top-10 list, counting a hit when "
             f"the user rated the title {C.RELEVANT_ANIME} or more. Low rating error does not automatically give "
             f"good top-10 lists; popularity is a strong ranking baseline on this data.")

# =================================================================== 12 clustering
if "clustering" in R:
    cl = R["clustering"]
    s = new()
    title(s, "Result 4: anime embeddings and clusters")
    s.shapes.add_picture(str(C.FIG_DIR / "anime_embedding_tsne.png"), Inches(0.5), Inches(1.35), height=Inches(5.5))
    stat(s, 7.7, 1.5, 2.6, f"{cl['genre_jaccard_within_spectral']:.2f}", "genre overlap inside a cluster", BLUE, 40)
    stat(s, 10.3, 1.5, 2.6, f"{cl['genre_jaccard_random_pairs']:.2f}", "genre overlap of random pairs", GREY, 40)
    stat(s, 7.7, 3.3, 2.6, f"{cl['genre_jaccard_title_vs_opposite']:.2f}", "title vs its 'opposite' title", ORANGE, 40)
    stat(s, 10.3, 3.3, 2.6, f"{cl['n_clusters']}", f"spectral clusters (median size {cl['size_median']:.0f})", INK, 40)
    text(s, 7.7, 5.3, 5.1, 1.5, "Clusters group anime with shared genres; Algorithm 1 picks titles whose genres "
         "differ more than a random title's.", size=15, color=INK2)
    notes(s, "Genre overlap is the Jaccard similarity of two titles' genre sets. Inside a cluster it is much higher "
             "than for random pairs, so the embeddings capture content. The 'opposite' title shares fewer genres "
             "than random, which is what Algorithm 1 aims for.")

# =================================================================== 13 hybrid filter
if "hybrid_filter" in R:
    hy = R["hybrid_filter"]
    s = new()
    title(s, "Result 5: the hybrid filter in action")
    keys = ["Autoencoder top-10 (no filter)", "Hybrid: Similar Anime", "Hybrid: Anime You May Like"]
    bar_chart(s, 0.5, 1.4, 6.2, 2.7, ["AE top-10", "Similar Anime", "You May Like"],
              [hy[k]["genre_jaccard_to_context"] for k in keys], [GREY, BLUE, BLUE],
              title_="Genre overlap with the user's context titles", number_format="0.000")
    bar_chart(s, 0.5, 4.2, 6.2, 2.7, ["AE top-10", "Similar Anime", "You May Like"],
              [hy[k]["Precision@10"] for k in keys], [GREY, BLUE, BLUE],
              title_="Precision@10 on held-out ratings", number_format="0.000")
    ex = pd.read_csv(C.OUT_DIR / "example_recommendations.csv")
    exs = ex[ex.list != "Anime You May Like"].head(8)
    exs = exs.assign(anime=exs.anime.str.slice(0, 34))[["list", "anime", "rating / predicted"]]
    exs.columns = ["List", "Anime", "Rating"]
    exs["Rating"] = [f"{v:g}" if float(v).is_integer() else f"{v:.2f}" for v in exs["Rating"]]
    exs["List"] = exs["List"].replace({"context (last 3 known ratings)": "Context"})
    text(s, 7.0, 1.45, 5.9, 0.4, f"Example: Kaggle user {R.get('example_user_kaggle_id')}", size=15, bold=True)
    table(s, 7.0, 1.95, 5.9, exs, col_w=[1.4, 3.6, 0.9], size=11, row_h=0.38)
    notes(s, "Top chart: the Similar list is closer in genre to what the user just rated, which is its purpose. "
             "Bottom chart: hit rate against held-out ratings. Right: a real example from the test set.")

# =================================================================== 14 discussion
s = new()
title(s, "What we learned")
items = [
    ("Pipeline works", "Every RikoNet component was implemented and produces sensible, genre-coherent lists."
     if "clustering" in R else "The rating-prediction part of RikoNet was implemented and evaluated on the paper's benchmark."),
    ("Headline not reproduced", f"MovieLens RMSE {fmt(e1o.RMSE)} vs 0.591 reported; {f['e1_vs_baseline']} "
                                f"a tuned baseline."),
    ("Reporting matters", "Unreported hyper-parameters, borrowed baseline numbers and table errors make "
                          "the claim hard to verify."),
    ("RMSE ≠ good lists", "Rating error and top-10 quality disagree; recommenders need ranking metrics too.")
    if "E4" in F else ("SELU confirmed", "Of four activations, SELU gave the lowest validation error, as the paper says."
                       if f.get("E3_movielens_best") == "selu" else "The paper's SELU claim did not hold on MovieLens."),
]
for i, (h_, b_) in enumerate(items):
    y = 1.6 + i * 1.3
    badge(s, 0.7, y + 0.1, i + 1)
    text(s, 1.45, y, 11, 0.5, h_, size=20, bold=True)
    text(s, 1.45, y + 0.5, 11, 0.7, b_, size=16, color=INK2)
notes(s, "Summarise the takeaways. Emphasise that a failed reproduction is still a valid scientific result if it "
         "is done carefully and transparently.")

# =================================================================== 15 conclusion
s = new(dark=True)
text(s, 0.8, 1.2, 11.5, 1.0, "Conclusion & future work", size=40, bold=True, color=WHITE, font=HEAD)
bullets(s, 0.8, 2.5, 11.5, 3.6, [
    "Simplified replication of RikoNet completed: exact benchmark data, alternative anime data.",
    f"Our RikoNet: RMSE {fmt(e1o.RMSE)} on MovieLens-100K; the paper's 0.591 was not reproduced.",
    "Next: ask the authors for code/data, try item-based autoencoders, run a small user study.",
], size=20, color=WHITE, gap=14)
text(s, 0.8, 6.2, 11.5, 0.5, "Thank you - questions?", size=22, bold=True, color="9FB8E8")
notes(s, "Close with the main finding and the next steps, then open the floor for questions.")

# =================================================================== 16 references
s = new()
title(s, "References")
bullets(s, 0.6, 1.5, 12.1, 5.3, [
    "Soni B., Thakuria D., Nath N., Das N., Boro B. RikoNet: A novel anime recommendation engine. "
    "Multimedia Tools and Applications 82:32329-32348 (2023). doi:10.1007/s11042-023-14710-9",
    "Kuchaiev O., Ginsburg B. Training deep autoencoders for collaborative filtering. arXiv:1708.01715 (2017).",
    "Zelnik-Manor L., Perona P. Self-tuning spectral clustering. NIPS 17 (2004).",
    "Klambauer G. et al. Self-normalizing neural networks. NIPS (2017).",
    "Harper F. M., Konstan J. A. The MovieLens datasets: history and context. ACM TiiS 5(4) (2015).",
    "CooperUnion. Anime Recommendations Database, Kaggle.",
], size=14)
notes(s, "Reference list.")

out = C.OUT_DIR / "RikoNet_Replication_Presentation.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
